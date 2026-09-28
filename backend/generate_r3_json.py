"""
Generate R3 JSON files from Prophet forecast Excel data.

Reads from: Emission_Papers/Shared_Data/Cambium_Powerplant/{gas}/{region}/harmonicmean/results/forecasted_results.xlsx
Writes to:   frontend/src/assets/r3data/{STATE}_{GAS}.json

Units (from Cambium/Prophet → display):
  CO2: kg/MWh → ton/MWh  (÷ 1000)
  CH4: g/MWh  → lb/MWh   (÷ 453.592)
  N2O: g/MWh  → lb/MWh   (÷ 453.592)
"""

import json
import os
import openpyxl

# ── Paths ──────────────────────────────────────────────────────────────────────
SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_ROOT = os.path.dirname(SCRIPT_DIR)
CAMBIUM_ROOT = os.path.join(
    os.path.dirname(PROJECT_ROOT),
    "Emission_Papers", "Shared_Data", "Cambium_Powerplant",
)
OUTPUT_DIR = os.path.join(PROJECT_ROOT, "frontend", "src", "assets", "r3data")

# ── Mappings ───────────────────────────────────────────────────────────────────
STATE_TO_REGION = {
    "GA": "SERTP",
    "CA": "CAISO",
    "NY": "NYISO",
    "WA": "NorthernGrid_West",
}

GASES = ["CO2", "CH4", "N2O"]

# Sheet-name scenario key → display name
SCENARIO_MAP = {
    "Midcase":      "Mid-case Scenario",
    "LowRE":        "Low Renewable Energy and Battery Costs",
    "HighRE":       "High Renewable Energy and Battery Costs",
    "HighDemand":   "High Demand Growth",
    "LowNG":        "Low Natural Gas Prices",
    "HighNG":       "High Natural Gas Prices",
    "LowRE_HighNG": "Low RE & Battery Costs + High Gas Prices",
    "HighRE_LowNG": "High RE & Battery Costs + Low Gas Prices",
}

# Ordered scenario keys (determines JSON key order)
SCENARIO_ORDER = list(SCENARIO_MAP.keys())

# Unit conversion divisors
UNIT_DIVISOR = {
    "CO2": 1000.0,      # kg/MWh → ton/MWh
    "CH4": 453.592,     # g/MWh  → lb/MWh
    "N2O": 453.592,     # g/MWh  → lb/MWh
}


def read_sheet(wb, sheet_name, value_col, divisor):
    """Read a sheet and return list of {date, value} dicts."""
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(min_row=2, values_only=True))  # skip header
    result = []
    for row in rows:
        dt = row[0]  # datetime
        raw_val = row[1] if len(row) > 1 else None
        if dt is None or raw_val is None:
            continue
        date_str = dt.strftime("%Y-%m-%d")
        value = round(raw_val / divisor, 6)
        result.append({"date": date_str, "value": value})
    return result


def generate_json(state, gas):
    """Generate one JSON file for a state/gas combination."""
    region = STATE_TO_REGION[state]
    xlsx_path = os.path.join(
        CAMBIUM_ROOT, gas, region, "harmonicmean", "results", "forecasted_results.xlsx"
    )

    if not os.path.exists(xlsx_path):
        print(f"  SKIP {state}_{gas}: {xlsx_path} not found")
        return False

    wb = openpyxl.load_workbook(xlsx_path, read_only=True)
    divisor = UNIT_DIVISOR[gas]

    scenarios = {}
    actual = {}

    for scenario_key in SCENARIO_ORDER:
        display_name = SCENARIO_MAP[scenario_key]
        forecast_sheet = f"hm_forecast_{scenario_key}_{gas}"
        actual_sheet = f"hm_actual_{scenario_key}_{gas}"

        if forecast_sheet not in wb.sheetnames:
            print(f"  WARNING: sheet {forecast_sheet} not found in {xlsx_path}")
            continue

        # Forecast data (312 monthly points, 2025-01 to 2050-12)
        scenarios[display_name] = read_sheet(wb, forecast_sheet, "yhat", divisor)

        # Actual data (72 data points — Cambium simulation values)
        if actual_sheet in wb.sheetnames:
            actual[display_name] = read_sheet(wb, actual_sheet, "y", divisor)

    wb.close()

    output = {"scenarios": scenarios, "actual": actual}
    out_path = os.path.join(OUTPUT_DIR, f"{state}_{gas}.json")
    with open(out_path, "w") as f:
        json.dump(output, f)

    # Summary
    first_scenario = list(scenarios.values())[0] if scenarios else []
    first_actual = list(actual.values())[0] if actual else []
    n_forecast = len(first_scenario)
    n_actual = len(first_actual)
    first_val = first_scenario[0]["value"] if first_scenario else "?"
    last_val = first_scenario[-1]["value"] if first_scenario else "?"
    print(f"  {state}_{gas}.json: {len(scenarios)} scenarios, "
          f"{n_forecast} forecast pts, {n_actual} actual pts, "
          f"first={first_val}, last={last_val}")
    return True


def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    print(f"Cambium data: {CAMBIUM_ROOT}")
    print(f"Output dir:   {OUTPUT_DIR}")
    print()

    for state in STATE_TO_REGION:
        print(f"[{state}]")
        for gas in GASES:
            generate_json(state, gas)
        print()

    print("Done!")


if __name__ == "__main__":
    main()
